"""Экспорт подтверждённой расшифровки и результатов проверки по ключам."""

import argparse
import json
from pathlib import Path

from .corrections import apply_corrections
from .evaluate import evaluate
from .recognize import relative, write_review


def text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        # Сохраняем ответы текстом, включая =, минусы, ведущие нули и запятые.
        cell.data_type = 's'
        cell.number_format = '@'
    return cell


def add_row(sheet, values):
    row = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        text_cell(sheet, row, column, value)


def workbook(data, metrics, destination):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ValueError('Установите экспорт: .venv/bin/python -m pip install -e ".[export]"') from None
    book = Workbook()
    students = book.active
    students.title = 'Ученики'
    details = book.create_sheet('Поля')
    errors = book.create_sheet('Ошибки OCR')
    replacements = book.create_sheet('Проверка замен')
    info = book.create_sheet('Сведения')
    add_row(students, ['Страница', 'Фамилия', 'Имя', 'Отчество', 'Класс'] +
            [f'Ответ {number}' for number in range(1, 20)] +
            ['Проверено полей', 'Всего полей', 'Статус расшифровки', 'Блок замены', 'Вариант', 'Баллы', 'Статус оценивания', 'Исходный скан',
             'Предварительный балл', 'Максимум', 'Демонстрационные ключи'])
    add_row(details, ['Страница', 'Поле', 'Текст OCR', 'Кандидат', 'Подтверждённое значение',
                      'Проверено вручную', 'Уверенность OCR', 'Изменено правилом', 'Фрагмент скана'])
    add_row(errors, ['Поле', 'Кандидат OCR', 'Ручная расшифровка'])
    add_row(replacements, ['Страница', 'Изображение блока', 'Статус', 'Применённые замены'])
    add_row(info, ['Параметр', 'Значение'])
    scores = book.create_sheet('Проверка ответов')
    add_row(scores, ['Страница', 'Вариант', 'Задание', 'Ответ', 'Источник ответа', 'Ключ', 'Балл', 'Максимум', 'Результат', 'ДЕМО'])
    for page in data['pages']:
        by_name = {field['name']: field for field in page['fields']}
        def confirmed(name):
            if name.startswith('answer_') and page.get('replacements_reviewed') is True:
                number = str(int(name[7:]))
                if number in page.get('replacement_answers', {}):
                    return page['replacement_answers'][number]
            field = by_name.get(name)
            return field['corrected_text'] if field and field['reviewed'] else None
        reviewed = sum(field['reviewed'] for field in page['fields'])
        complete = len(by_name) == 23 and reviewed == 23 and page['status'] == 'prepared'
        grading = page.get('grading', {})
        replacement_status = 'Проверен' if page.get('replacements_reviewed') is True else 'Не проверен'
        add_row(students, [page['page']] + [confirmed(name) for name in
                ['surname', 'first_name', 'patronymic', 'class']] +
                [confirmed(f'answer_{number:02}') for number in range(1, 20)] +
                [reviewed, len(by_name), 'Проверена вручную' if complete else 'Требуется проверка',
                 replacement_status, grading.get('variant'), grading.get('total'),
                 grading.get('status', 'Не выполнено: нужны ключи, варианты и проверка замен'), page.get('original'),
                 grading.get('provisional'), grading.get('maximum'), 'Да' if grading.get('demo') else 'Нет'])
        add_row(replacements, [page['page'], page.get('replacements'), replacement_status,
                              json.dumps(page.get('replacement_answers', {}), ensure_ascii=False) if page.get('replacements_reviewed') else ''])
        for task in grading.get('tasks', []):
            add_row(scores, [page['page'], grading['variant'], task['number'], task['answer'], task['source'],
                             ' / '.join(task['accepted']), task['points'], task['maximum'], task['status'],
                             'Да' if grading['demo'] else 'Нет'])
        for field in page['fields']:
            add_row(details, [page['page'], field['label'], field['raw_text'], field['candidate'],
                              field['corrected_text'] if field['reviewed'] else None,
                              'Да' if field['reviewed'] else 'Нет', field['confidence'],
                              ', '.join(item['rule'] for item in field.get('adjustments', [])), field['original']])
    for mismatch in metrics['mismatches']:
        add_row(errors, [mismatch['id'], mismatch['ocr'], mismatch['expected']])
    for key, value in [
        ('Содержание', 'Расшифровка и проверка кратких ответов 1–19; развёрнутые решения не оцениваются'),
        ('Ключи', data.get('answer_key', {}).get('title', 'Не загружены')),
        ('ДЕМО', 'Случайные ответы! Баллы не являются результатами ОГЭ.' if data.get('answer_key', {}).get('demo') else 'Нет'),
        ('Страниц', len(data['pages'])),
        ('Подтверждено полей', metrics['reviewed_fields']),
        ('Точно распознано до ручных правок', metrics['matching_fields']),
        ('Доля точных совпадений', metrics['exact_match_rate']),
        ('Ограничение метрики', 'OCR сравнивается с ручными правками. Это не точность выставления баллов.'),
        ('Блок замены', 'Учитывается только после отдельной ручной проверки блока'),
        ('Пустые значения', 'Проверенный пустой ответ и непроверенное поле различаются в листе Поля'),
        ('Ручные правки', 'Сохранены дословно; регистр и смешение алфавитов не исправлялись'),
        ('Источник', data['source_manifest']),
    ]:
        add_row(info, [key, value])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='254A70')
        for column in sheet.columns:
            width = min(55, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    book.save(destination)


def export(recognition_path, corrections_path, output):
    recognition_path, corrections_path, output = (path.resolve() for path in
                                                  (recognition_path, corrections_path, output))
    recognition = json.loads(recognition_path.read_text(encoding='utf-8'))
    corrections = json.loads(corrections_path.read_text(encoding='utf-8'))
    if recognition.get('run_status') != 'complete':
        raise ValueError('Сначала завершите запуск OCR')
    data = apply_corrections(recognition, corrections)
    metrics = evaluate(recognition, corrections)
    if output.exists() and any(output.iterdir()):
        raise ValueError('Папка экспорта должна быть пустой; укажите новую --output')
    output.mkdir(parents=True, exist_ok=True)
    # Ссылки должны работать из новой папки, не из папки исходного OCR.
    for page in data['pages']:
        for key in ('original', 'replacements'):
            if key in page:
                page[key] = relative(recognition_path.parent / page[key], output)
        for field in page['fields']:
            for key in ('original', 'cleaned'):
                field[key] = relative(recognition_path.parent / field[key], output)
    data['corrections_source'] = str(corrections_path)
    workbook(data, metrics, output / 'students.xlsx')
    (output / 'confirmed.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / 'metrics.json').write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding='utf-8')
    write_review(output, data)
    print(f'Подтверждено полей: {metrics["reviewed_fields"]}; таблица: {output / "students.xlsx"}')
    return data, metrics


def main():
    parser = argparse.ArgumentParser(description='Экспорт ручных правок в Excel и JSON')
    parser.add_argument('recognition', type=Path)
    parser.add_argument('corrections', type=Path)
    parser.add_argument('--output', type=Path, default=Path('output/results'))
    args = parser.parse_args()
    try:
        export(args.recognition, args.corrections, args.output)
    except (OSError, ValueError, KeyError, TypeError) as exc:
        parser.exit(1, f'Ошибка: {exc}\n')


if __name__ == '__main__':
    main()
